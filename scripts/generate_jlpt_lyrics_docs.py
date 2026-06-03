import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def safe_filename(text):
    text = re.sub(r'[<>:"/\\\\|?*]+', "_", text)
    text = re.sub(r"\s+", "_", text).strip("._ ")
    return text or "jlpt-lyrics"


def set_run_font(run, size=None, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    if size:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, width in enumerate(widths):
            cell = row.cells[i]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def fill_cell(cell, text, size=9.0, bold=False, color=None, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold, color=color)


def add_header(row, labels):
    for cell, label in zip(row.cells, labels):
        set_cell_shading(cell, "E8EEF5")
        fill_cell(cell, label, 9.2, True, "0B2545", WD_ALIGN_PARAGRAPH.CENTER)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    r = p.add_run(text)
    set_run_font(r, {1: 16, 2: 13, 3: 12}[level], True, {1: "2E74B5", 2: "2E74B5", 3: "1F4D78"}[level])


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, 9.4, False, "555555")


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
    return doc


def title_block(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, 21, True, "0B2545")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run(subtitle)
    set_run_font(r, 10, False, "555555")


def add_table(doc, labels, rows, widths, size=8.8):
    table = doc.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    add_header(table.rows[0], labels)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, len(row) - 1) and len(str(text)) <= 8 else None
            fill_cell(cells[i], text, size=size, align=align)
    set_table_widths(table, widths)
    return table


def build_handout(data, out_dir):
    title = data["song_title"]
    title_cn = data.get("song_title_cn")
    artist = data.get("artist", "")
    target_level = data.get("target_level", "").upper()
    display = f"{title} / {title_cn}" if title_cn else title

    doc = setup_doc()
    level_text = f"{target_level} " if target_level else ""
    title_block(doc, display, f"{artist} | {level_text}JLPT 歌词学习讲义（{data.get('source_note', '基于用户提供文本整理')}）")
    add_note(doc, "阅读方法：左栏为日语原文学习版，汉字后用括号标注假名；右栏为便于理解的自然中文译文。")

    add_heading(doc, "一、歌词对照", 1)
    rows = [[item["ja"], item["zh"]] for item in data.get("lyrics", [])]
    add_table(doc, ["日语歌词原文（汉字假名标注）", "中文译文"], rows, [4700, 4660], 9.0)

    add_heading(doc, "二、生词表", 1)
    vocab_rows = [
        [v.get("word", ""), v.get("kana", ""), v.get("meaning", ""), v.get("pos", ""), v.get("jlpt", ""), v.get("usage", "")]
        for v in data.get("vocabulary", [])
    ]
    add_table(doc, ["单词", "假名", "中文意思", "词性", "JLPT", "歌词中的用法"], vocab_rows, [1350, 1500, 1500, 1200, 700, 3110], 8.2)

    add_heading(doc, "三、语法点讲解", 1)
    grammar_rows = [
        [
            g.get("pattern", ""),
            g.get("jlpt", ""),
            " ".join(x for x in [g.get("explanation", ""), g.get("lyric_context", "")] if x),
            g.get("example", ""),
        ]
        for g in data.get("grammar", [])
    ]
    add_table(doc, ["语法", "JLPT", "讲解", "例句"], grammar_rows, [1650, 900, 4420, 2390], 8.2)

    tips = data.get("study_tips", [])
    if tips:
        add_heading(doc, "四、学习提示", 1)
        for tip in tips:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(tip)
            set_run_font(r, 10)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run(f"JLPT 歌词学习讲义 | {title}"), 8.5, False, "777777")

    level_part = f"_{safe_filename(target_level)}" if target_level else ""
    path = out_dir / f"{safe_filename(title)}_{safe_filename(artist)}{level_part}_JLPT歌词学习讲义.docx"
    doc.save(path)
    return path


def build_drills(data, out_dir):
    title = data["song_title"]
    artist = data.get("artist", "")
    target_level = data.get("target_level", "").upper()
    doc = setup_doc()
    level_text = f"{target_level} " if target_level else ""
    title_block(doc, f"{title} / {level_text}JLPT 应试训练", f"{artist} | 不做歌词默写或听力模拟，重点训练识别、理解与运用")
    add_note(doc, "设计说明：本卷把歌词作为语境材料，题型对应 JLPT 的文字词汇、语法、读解和信息检索。")

    for section in data.get("drill_sections", []):
        ensure_no_listening(section)
        add_heading(doc, section["title"], 1)
        for task in section.get("tasks", []):
            ensure_no_listening(task)
            add_heading(doc, task["title"], 2)
            if task.get("instruction"):
                add_note(doc, task["instruction"])
            cols = task.get("columns", ["No.", "题干", "选项", "答"])
            widths = task.get("widths") or default_widths(len(cols))
            add_table(doc, cols, task.get("rows", []), widths, task.get("font_size", 8.4))

    answers = data.get("answers", [])
    if answers:
        add_heading(doc, "答案与解析", 1)
        for block in answers:
            add_heading(doc, block.get("title", "答案"), 2)
            for item in block.get("items", []):
                p = doc.add_paragraph()
                r = p.add_run(item)
                set_run_font(r, 9.4)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run(f"JLPT 分级应试训练 | {title}"), 8.5, False, "777777")

    level_part = f"_{safe_filename(target_level)}" if target_level else ""
    path = out_dir / f"{safe_filename(title)}_{safe_filename(artist)}{level_part}_JLPT应试训练.docx"
    doc.save(path)
    return path


def ensure_no_listening(obj):
    text = json.dumps(obj, ensure_ascii=False).lower()
    banned = ["听力", "聴解", "listening", "teacher-read", "教师朗读", "朗读脚本"]
    hit = [word for word in banned if word.lower() in text]
    if hit:
        raise ValueError(f"Listening-style tasks are disabled for this skill. Remove: {', '.join(hit)}")


def validate_target_level(data):
    level = data.get("target_level")
    if not level:
        raise ValueError("input JSON must include target_level: N5, N4, N3, N2, or N1")
    level = str(level).upper()
    if level not in {"N5", "N4", "N3", "N2", "N1"}:
        raise ValueError("target_level must be one of: N5, N4, N3, N2, N1")
    data["target_level"] = level


def default_widths(col_count):
    if col_count == 2:
        return [900, 8460]
    if col_count == 3:
        return [700, 3800, 4860]
    if col_count == 4:
        return [650, 3600, 4200, 910]
    return [9360 // col_count] * col_count


def main():
    parser = argparse.ArgumentParser(description="Generate JLPT lyric lesson DOCX files from a JSON payload.")
    parser.add_argument("input_json", type=Path)
    parser.add_argument("--out-dir", type=Path, default=Path("outputs"))
    parser.add_argument("--only", choices=["handout", "drills", "both"], default="both")
    args = parser.parse_args()

    data = json.loads(args.input_json.read_text(encoding="utf-8"))
    validate_target_level(data)
    args.out_dir.mkdir(parents=True, exist_ok=True)

    outputs = []
    if args.only in ("handout", "both"):
        outputs.append(build_handout(data, args.out_dir))
    if args.only in ("drills", "both"):
        outputs.append(build_drills(data, args.out_dir))

    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
